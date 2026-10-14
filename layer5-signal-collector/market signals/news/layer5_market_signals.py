#!/usr/bin/env python3
"""Layer 5 - Market signals: daily headline collector for all sectors.

Every sector folder next to this script holds:
  <sector>-rss-feeds.csv            the feed definitions (3 always-on + 35 rotating, 5 per weekday)
  <sector>-market-signals.xlsx      one workbook; the script adds one sheet per day (named YYYY-MM-DD)

What one run does (see METHOD.md):
  1. picks today's feeds per sector = always-on feeds + the feeds assigned to today's weekday
  2. fetches each Google News RSS search feed (when:1d = last 24 hours), spaced out with random gaps
  3. keeps items published in the last 24 hours, drops junk (quote pages, option contracts), removes duplicates
  4. appends the new headlines to today's sheet in the sector workbook (re-runs never create duplicates)
  5. logs every feed to run-log.csv (status, items, stale items, whether the date window held, cap hit)
  6. resolves each Google News redirect link to the real publisher URL (gnews_resolve.py, one folder up) and
     stores it in the article_url column of today's sheet; logs the counts to resolve-log.csv.
     Links that cannot be resolved stay blank (source_link still works as a fallback).

Usage:
  python layer5_market_signals.py --dry-run                     show today's feeds for every sector, fetch nothing
  python layer5_market_signals.py --sector space                one sector
  python layer5_market_signals.py --sector space,defense        several sectors
  python layer5_market_signals.py                               all 26 sectors
Options: --day Monday..Sunday (use another weekday's feeds), --gap-min/--gap-max seconds between requests,
         --cooldown seconds to wait after a 503/429, --limit-feeds N (testing),
         --no-resolve (skip step 6), --resolve-only (only step 6, on today's existing sheets),
         --resolve-budget-min N (stop resolving after N minutes, default 45), --resolve-workers N (default 4).
"""
import argparse, csv, html, os, random, re, sys, time, urllib.error, urllib.parse, urllib.request
import email.utils
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

try:
    from zoneinfo import ZoneInfo
    ET = ZoneInfo("America/New_York")
except Exception:                      # no tz database: fall back to a fixed offset
    ET = timezone(timedelta(hours=-4))

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))          # gnews_resolve.py sits in "market signals/"
import gnews_resolve
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
FEED_URL = "https://news.google.com/rss/search?q={q}&hl=en-US&gl=US&ceid=US:en"
HEADER = ["published_et", "headline", "outlet", "source_link", "outlet_site", "feed_ids", "sector", "collected_et",
          "article_url"]
WIDTHS = [18, 100, 26, 70, 30, 30, 34, 18, 90]
JUNK = re.compile(r"stock price, news, quote|quote and history|\b\w+\d{6}[CP]\d{5,8}\b|\b(call|put) \(|price prediction|stock forecasts?\b|share price history", re.I)
RESOLVE_LOG_FIELDS = ["run_at_et", "sector", "links", "resolved_offline", "resolved_online", "failed",
                      "skipped_budget", "note"]
LOG_FIELDS = ["run_at_et", "sector", "feed_id", "status", "http", "items", "in_window", "stale_dropped", "junk", "new_rows",
              "oldest_item_et", "window_ok", "capped_100", "note"]


def norm(t):
    t = re.sub(r"\s+-\s+[^-]{2,40}$", "", t.lower())      # drop " - Outlet" suffix
    return re.sub(r"[^a-z0-9 ]", "", t).strip()[:90]


def fetch(query, timeout=40):
    url = FEED_URL.format(q=urllib.parse.quote(query, safe=":()"))
    try:
        with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": UA}), timeout=timeout) as r:
            return r.status, r.read().decode("utf-8", "ignore")
    except urllib.error.HTTPError as e:
        return e.code, ""
    except Exception as e:
        return 0, str(e)[:120]


def parse(xml):
    out = []
    for it in re.findall(r"<item>.*?</item>", xml, re.S):
        m = re.search(r"<title>(.*?)</title>", it, re.S)
        if not m:
            continue
        title = html.unescape(re.sub(r"<!\[CDATA\[|\]\]>", "", m.group(1))).strip()
        link = (re.search(r"<link>(.*?)</link>", it, re.S) or [None, ""])[1].strip()
        pd = re.search(r"<pubDate>(.*?)</pubDate>", it)
        try:
            dt = email.utils.parsedate_to_datetime(pd.group(1)).astimezone(ET)
        except Exception:
            continue
        s = re.search(r'<source[^>]*url="([^"]*)"[^>]*>(.*?)</source>', it, re.S)
        site, outlet = (html.unescape(s.group(1)), html.unescape(s.group(2)).strip()) if s else ("", "")
        head = re.sub(r"\s+-\s+" + re.escape(outlet) + r"$", "", title) if outlet else title
        out.append({"headline": head, "outlet": outlet, "link": link, "site": site, "dt": dt})
    return out


def read_feeds(sector_dir, slug, weekday):
    path = os.path.join(sector_dir, slug + "-rss-feeds.csv")
    if not os.path.exists(path):
        return None
    rows = list(csv.DictReader(open(path, encoding="utf-8-sig")))
    return [r for r in rows if r["run_day"] in ("Every day", weekday)]


def open_book(path):
    if os.path.exists(path):
        return openpyxl.load_workbook(path)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


def get_sheet(wb, name):
    if name in wb.sheetnames:
        ws = wb[name]
        n = len(HEADER)
        if ws.cell(1, n).value != HEADER[-1]:            # sheet made before the article_url column existed
            c = ws.cell(1, n)
            c.value = HEADER[-1]
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
            ws.column_dimensions[get_column_letter(n)].width = WIDTHS[-1]
        return ws, False
    ws = wb.create_sheet(name, 0)             # newest day first
    ws.append(HEADER)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws, True


def write_sector(sector_dir, slug, sector_name, today, day_items, now_et):
    """day_items: list of dict(headline,outlet,link,site,dt,feed_ids:set). Returns (new_rows, updated_rows, saved_path)."""
    path = os.path.join(sector_dir, f"{slug}-market-signals.xlsx")
    wb = open_book(path)
    ws, _ = get_sheet(wb, today)
    index = {}                                 # dedupe key -> row number
    for r in range(2, ws.max_row + 1):
        h, link = ws.cell(r, 2).value, ws.cell(r, 4).value
        if h:
            index[norm(str(h))] = r
        if link:
            index["L:" + str(link)] = r
    new = upd = 0
    for it in sorted(day_items, key=lambda x: x["dt"], reverse=True):
        k, kl = norm(it["headline"]), "L:" + it["link"]
        row = index.get(k) or index.get(kl)
        feeds = ", ".join(sorted(it["feed_ids"]))
        if row:
            have = {x.strip() for x in str(ws.cell(row, 6).value or "").split(",") if x.strip()}
            if not it["feed_ids"] <= have:
                ws.cell(row, 6).value = ", ".join(sorted(have | it["feed_ids"]))
                upd += 1
            continue
        ws.append([it["dt"].strftime("%Y-%m-%d %H:%M"), it["headline"], it["outlet"], it["link"], it["site"], feeds,
                   sector_name, now_et.strftime("%Y-%m-%d %H:%M"), ""])
        index[k] = index[kl] = ws.max_row
        new += 1
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}{ws.max_row}"
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 2).alignment = Alignment(wrap_text=False)
    try:
        wb.save(path)
    except PermissionError:                   # workbook open in Excel
        path = os.path.join(sector_dir, f"{slug}-market-signals.PENDING-{today}.xlsx")
        wb.save(path)
    return new, upd, path


def resolve_phase(sector_paths, today, now, workers, budget_min):
    """sector_paths: {sector: workbook path}. Fills article_url on today's sheet for rows that are still blank."""
    import threading
    deadline = time.time() + budget_min * 60
    stop = {"flag": False, "streak": 0}
    lock = threading.Lock()
    cache = {}                                          # google link -> (url, status)
    n_col = len(HEADER)

    def work(link):
        if stop["flag"] or time.time() > deadline:
            return link, ("", "skipped")
        time.sleep(random.uniform(0.05, 0.25))
        url, st = gnews_resolve.resolve(link)
        with lock:
            if url:
                stop["streak"] = 0
            elif st.startswith("http 429") or st.startswith("http 503"):
                stop["streak"] += 1
                if stop["streak"] >= 25:
                    stop["flag"] = True
            else:
                stop["streak"] = 0
        return link, (url, st)

    # collect the blank links first so a link shared by several sectors is resolved once
    todo = {}
    books = {}
    for s, path in sector_paths.items():
        if not os.path.exists(path):
            continue
        wb = openpyxl.load_workbook(path)
        if today not in wb.sheetnames:
            continue
        ws, _ = get_sheet(wb, today)
        books[s] = (wb, path)
        for r in range(2, ws.max_row + 1):
            link, cur = ws.cell(r, 4).value, ws.cell(r, n_col).value
            if link and not cur and gnews_resolve.article_id(str(link)):
                todo.setdefault(str(link), None)
    print(f"resolve: {len(todo)} unresolved links across {len(books)} sectors "
          f"({workers} workers, budget {budget_min} min)")
    t0 = time.time()
    with ThreadPoolExecutor(max_workers=workers) as ex:
        for link, res in ex.map(work, list(todo)):
            cache[link] = res
    print(f"resolve: finished in {(time.time() - t0) / 60:.1f} min"
          + ("  [STOPPED early: repeated 429/503]" if stop["flag"] else ""))

    log_path = os.path.join(HERE, "resolve-log.csv")
    new_log = not os.path.exists(log_path)
    with open(log_path, "a", newline="", encoding="utf-8-sig") as lf:
        lw = csv.DictWriter(lf, fieldnames=RESOLVE_LOG_FIELDS)
        if new_log:
            lw.writeheader()
        for s, (wb, path) in books.items():
            ws = wb[today]
            n = off = onl = fail = skip = 0
            for r in range(2, ws.max_row + 1):
                link, cur = ws.cell(r, 4).value, ws.cell(r, n_col).value
                if cur or not link or str(link) not in cache:
                    continue
                url, st = cache[str(link)]
                n += 1
                if url:
                    ws.cell(r, n_col).value = url
                    off += st == "offline"
                    onl += st == "online"
                elif st == "skipped":
                    skip += 1
                else:
                    fail += 1
            try:
                wb.save(path)
            except PermissionError:
                print(f"  {s}: workbook is open elsewhere, article_url not saved")
                continue
            lw.writerow({"run_at_et": now.strftime("%Y-%m-%d %H:%M"), "sector": s, "links": n,
                         "resolved_offline": off, "resolved_online": onl, "failed": fail, "skipped_budget": skip,
                         "note": "stopped early (429/503)" if stop["flag"] else ""})
            print(f"  [{s}] links {n}: resolved {off + onl} (offline {off}), failed {fail}, skipped {skip}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sector", help="sector folder name(s), comma separated; default all")
    ap.add_argument("--day", choices=DAYS, help="use this weekday's rotating feeds (default: today)")
    ap.add_argument("--dry-run", action="store_true", help="list the feeds that would run and exit")
    ap.add_argument("--gap-min", type=float, default=3.0)
    ap.add_argument("--gap-max", type=float, default=7.0)
    ap.add_argument("--cooldown", type=int, default=600, help="seconds to wait after a 503/429 before one retry")
    ap.add_argument("--limit-feeds", type=int, default=0, help="only the first N feeds per sector (testing)")
    ap.add_argument("--no-resolve", action="store_true", help="skip resolving Google News links to article URLs")
    ap.add_argument("--resolve-only", action="store_true", help="only resolve links on today's existing sheets")
    ap.add_argument("--resolve-budget-min", type=float, default=45.0, help="stop resolving after this many minutes")
    ap.add_argument("--resolve-workers", type=int, default=4)
    a = ap.parse_args()

    now = datetime.now(ET)
    weekday = a.day or DAYS[now.weekday()]
    today = now.strftime("%Y-%m-%d")
    cutoff = now - timedelta(hours=24)
    sectors = sorted(d for d in os.listdir(HERE) if os.path.isdir(os.path.join(HERE, d)) and not d.startswith((".", "_")))
    if a.sector:
        want = [s.strip() for s in a.sector.split(",")]
        bad = [s for s in want if s not in sectors]
        if bad:
            sys.exit(f"unknown sector folder(s): {bad}\navailable: {sectors}")
        sectors = want

    if a.resolve_only:
        paths = {s: os.path.join(HERE, s, f"{s}-market-signals.xlsx") for s in sectors}
        resolve_phase(paths, today, now, a.resolve_workers, a.resolve_budget_min)
        return

    plan = {}
    for s in sectors:
        feeds = read_feeds(os.path.join(HERE, s), s, weekday)
        if feeds is None:
            print(f"[skip] {s}: no {s}-rss-feeds.csv")
            continue
        plan[s] = feeds[: a.limit_feeds] if a.limit_feeds else feeds
    total = sum(len(v) for v in plan.values())
    print(f"{today} {now.strftime('%H:%M')} ET | weekday feeds: {weekday} | sectors: {len(plan)} | feeds to run: {total}")
    if a.dry_run:
        for s, feeds in plan.items():
            print(f"\n{s}  ({len(feeds)} feeds)")
            for f in feeds:
                print(f"  {f['feed_id']:<46} {f['feed_type']:<10} {f['label']}")
        return

    log_path = os.path.join(HERE, "run-log.csv")
    new_log = not os.path.exists(log_path)
    logf = open(log_path, "a", newline="", encoding="utf-8-sig")
    lw = csv.DictWriter(logf, fieldnames=LOG_FIELDS)
    if new_log:
        lw.writeheader()

    blocked_in_a_row = 0
    saved_paths = {}
    grand_new = 0
    first_request = True
    for s, feeds in plan.items():
        cfg_name = s
        day_items = {}                         # norm headline -> item
        for f in feeds:
            if blocked_in_a_row >= 2:
                break
            if not first_request:
                time.sleep(random.uniform(a.gap_min, a.gap_max))
            first_request = False
            q = f["query"] if "when:" in f["query"] else f["query"] + " when:1d"
            code, body = fetch(q)
            note = ""
            if code in (429, 503):
                print(f"  {s}: HTTP {code} on {f['feed_id']} - waiting {a.cooldown}s then one retry")
                time.sleep(a.cooldown)
                code, body = fetch(q)
                note = "retried after block"
            rec = {"run_at_et": now.strftime("%Y-%m-%d %H:%M"), "sector": s, "feed_id": f["feed_id"], "http": code}
            if code != 200:
                blocked_in_a_row += 1 if code in (429, 503) else 0
                rec.update(status="FAILED", items=0, in_window=0, stale_dropped=0, junk=0, new_rows=0, window_ok="", capped_100="",
                           oldest_item_et="", note=(note + " " + body[:80]).strip())
                lw.writerow(rec); logf.flush()
                print(f"  {f['feed_id']:<46} FAILED http={code}")
                continue
            blocked_in_a_row = 0
            items = parse(body)
            fresh = [i for i in items if i["dt"] >= cutoff]
            stale = len(items) - len(fresh)
            clean = [i for i in fresh if not JUNK.search(i["headline"])]
            junk = len(fresh) - len(clean)
            n_before = len(day_items)
            for i in clean:
                k = norm(i["headline"])
                if k in day_items:
                    day_items[k]["feed_ids"].add(f["feed_id"])
                else:
                    i["feed_ids"] = {f["feed_id"]}
                    day_items[k] = i
            oldest = min((i["dt"] for i in items), default=None)
            rec.update(status="ok", items=len(items), in_window=len(fresh), stale_dropped=stale, junk=junk,
                       new_rows=len(day_items) - n_before, window_ok=("yes" if stale == 0 else "NO"),
                       capped_100=("yes" if len(items) >= 100 else ""), oldest_item_et=oldest.strftime("%Y-%m-%d %H:%M") if oldest else "",
                       note=note)
            lw.writerow(rec); logf.flush()
            flag = ("  [window ignored: %d stale dropped]" % stale if stale else "") + ("  [100-item cap]" if len(items) >= 100 else "")
            print(f"  {f['feed_id']:<46} items {len(items):3d} | last24h {len(fresh):3d} | junk {junk:2d} | new {len(day_items) - n_before:3d}{flag}")
        sector_dir = os.path.join(HERE, s)
        sector_name = cfg_name
        new, upd, path = write_sector(sector_dir, s, sector_name, today, list(day_items.values()), now)
        grand_new += new
        saved_paths[s] = path
        print(f"[{s}] sheet {today}: +{new} new headlines, {upd} existing rows got extra feed ids -> {os.path.basename(path)}")
        if blocked_in_a_row >= 2:
            print("STOPPED: two blocked feeds in a row (Google rate limit). Remaining sectors skipped; see run-log.csv.")
            break
    logf.close()
    print(f"done: {grand_new} new headlines written. log: run-log.csv")
    if not a.no_resolve:
        resolve_phase(saved_paths, today, now, a.resolve_workers, a.resolve_budget_min)


if __name__ == "__main__":
    main()

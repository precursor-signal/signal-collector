#!/usr/bin/env python3
"""Layer 5 - Market signals: daily market movers (top 100, any sector).

What it does, each run:
  1. Reads the day's biggest gainers and losers from Yahoo Finance (100 + up to 100 rows).
     If Yahoo fails or returns too few rows it falls back to the Nasdaq API, then to StockAnalysis, and
     merges what they return (each row keeps the source it came from).
  2. Ranks all rows by size of the % move (up or down) and keeps the top N (default 100), regardless of sector.
  3. Marks each mover's sector:
       layer5_sector  - the Layer 5 sector(s) the ticker was researched under (layer5-ticker-sectors.csv)
       yahoo_sector / yahoo_industry - Yahoo's own classification, looked up for names not in Layer 5
  4. Looks up catalyst hints: for each mover, the newest Google News headlines (last 2 days) that name the company
     or its ticker, with the real article URL (gnews_resolve.py, one folder up). Up to 3 per mover go in the
     catalyst_hint_1..3 columns as "published_et | outlet | headline | url". A hint is only a lead for the reader:
     it is not verified as the cause of the move.
  5. Writes one sheet per market session (named YYYY-MM-DD) into market-movers.xlsx in this folder.
     A re-run for the same session replaces that sheet; sheets from other days are never touched.

Usage:
  python layer5_market_movers.py                      normal run
  python layer5_market_movers.py --dry-run            fetch and print, write nothing
  python layer5_market_movers.py --top 50             keep 50 instead of 100
  python layer5_market_movers.py --source nasdaq      start at a different source (test the fallbacks)
  python layer5_market_movers.py --no-profile         skip the Yahoo sector lookup for unmatched names
  python layer5_market_movers.py --no-hints           skip the catalyst-hint lookup

Needs: python 3, openpyxl.  Run after the US close for the closing movers (weekends return Friday's close).
"""
import argparse, csv, email.utils, html, http.cookiejar, json, os, random, re, sys, time, urllib.error, urllib.parse
import urllib.request
from datetime import datetime, timedelta, timezone

try:
    from zoneinfo import ZoneInfo
    ET = ZoneInfo("America/New_York")
except Exception:                      # no tz database: fall back to a fixed offset
    ET = timezone(timedelta(hours=-4))

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))          # gnews_resolve.py sits in "market signals/"
import gnews_resolve
XLSX = os.path.join(HERE, "market-movers.xlsx")
SECTOR_MAP = os.path.join(HERE, "layer5-ticker-sectors.csv")
LOG = os.path.join(HERE, "movers-run-log.csv")
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
      "Chrome/124.0 Safari/537.36")
HEADER = ["rank", "symbol", "name", "direction", "pct_change", "price", "change", "volume", "avg_volume", "market_cap",
          "layer5_sector", "yahoo_sector", "yahoo_industry", "sector_source", "source", "session_date", "collected_et",
          "catalyst_hint_1", "catalyst_hint_2", "catalyst_hint_3"]
LOG_FIELDS = ["run_at_et", "session_date", "source", "list", "http", "rows", "note"]
SOURCE_ORDER = ["yahoo", "nasdaq", "stockanalysis"]
NEWS_URL = "https://news.google.com/rss/search?q={q}&hl=en-US&gl=US&ceid=US:en"
JUNK = re.compile(r"stock price, news, quote|quote and history|\b\w+\d{6}[CP]\d{5,8}\b|\b(call|put) \(|price prediction"
                  r"|stock forecasts?\b|share price history|top (gainers|losers)|stocks? (making|moving)|biggest movers",
                  re.I)
SUFFIX = re.compile(r"[,.]?\s+(inc\.?|incorporated|corp\.?|corporation|co\.?|company|ltd\.?|limited|plc|llc|lp|n\.?v\.?"
                    r"|s\.?a\.?|ag|se|holdings?|group|technologies|class [a-c]|common stock|ordinary shares?"
                    r"|american depositary shares?|ads|etf|trust)\b.*$", re.I)


# ----------------------------------------------------------------------------------------------- helpers
def http_get(url, timeout=30):
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept-Language": "en-US,en;q=0.9",
                                               "Accept": "application/json,text/html,*/*"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, r.read(5_000_000).decode("utf-8", "ignore")
    except urllib.error.HTTPError as e:
        return e.code, ""
    except Exception as e:
        return 0, str(e)[:100]


def text(s):
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", s))).strip()


def num(s):
    """'15.146M' -> 15146000.0, '-56.71%' -> -56.71, '1,234' -> 1234.0, '--' -> None."""
    if s is None:
        return None
    s = str(s).strip().replace(",", "").replace("$", "").replace("%", "").replace("+", "")
    m = re.fullmatch(r"(-?\d+(?:\.\d+)?)\s*([KMBT]?)", s, re.I)
    if not m:
        return None
    return float(m.group(1)) * {"": 1, "K": 1e3, "M": 1e6, "B": 1e9, "T": 1e12}[m.group(2).upper()]


def norm(sym):
    return sym.strip().upper().replace(".", "-")


def log_rows(rows):
    new = not os.path.exists(LOG)
    try:
        with open(LOG, "a", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            if new:
                w.writerow(LOG_FIELDS)
            w.writerows(rows)
    except PermissionError:
        print("  (run-log.csv is open elsewhere; log lines not written)")


# ----------------------------------------------------------------------------------------------- sources
# Each fetcher returns (rows, session_date_or_None, log_lines). A row is a dict with the HEADER fields it knows.
def row(sym, name, direction, pct, price=None, chg=None, vol=None, avg=None, cap=None, source=""):
    return {"symbol": norm(sym), "name": name, "direction": direction, "pct_change": pct, "price": price,
            "change": chg, "volume": vol, "avg_volume": avg, "market_cap": cap, "source": source}


def src_yahoo(now):
    out, lines, sess = [], [], None
    for kind, path, direction in (("gainers", "gainers", "Gainer"), ("losers", "losers", "Loser")):
        url = f"https://finance.yahoo.com/markets/stocks/{path}/?count=100"
        code, h = http_get(url)
        got = []
        if code == 200:
            for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", h, re.S):
                c = [text(x) for x in re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)]
                if len(c) < 9 or not re.fullmatch(r"[A-Z0-9.\-]{1,10}", c[0]):
                    continue
                pct = num(c[5])
                if pct is None:
                    continue
                got.append(row(c[0], c[1], direction, pct, num(c[3].split(" ")[0]), num(c[4]), num(c[6]), num(c[7]),
                               num(c[8]), "yahoo"))
        out += got
        lines.append(["yahoo", kind, code, len(got), "" if got else "no rows parsed"])
        time.sleep(random.uniform(2, 4))
    # session date from SPY's last trading timestamp
    code, b = http_get("https://query1.finance.yahoo.com/v8/finance/chart/SPY?range=5d&interval=1d")
    if code == 200:
        try:
            ts = json.loads(b)["chart"]["result"][0]["meta"]["regularMarketTime"]
            sess = datetime.fromtimestamp(ts, ET).strftime("%Y-%m-%d")
        except Exception:
            pass
    return out, sess, lines


def src_nasdaq(now):
    out, lines, sess = [], [], None
    for ex in ("NASDAQ", "NYSE", "AMEX"):
        code, b = http_get(f"https://api.nasdaq.com/api/marketmovers?exchange={ex}")
        got = []
        if code == 200:
            try:
                st = json.loads(b)["data"]["STOCKS"]
                for key, direction in (("MostAdvanced", "Gainer"), ("MostDeclined", "Loser")):
                    blk = st.get(key) or {}
                    if not sess:
                        m = re.search(r"([A-Z][a-z]{2} \d{1,2}, \d{4})", blk.get("dataAsOf") or "")
                        if m:
                            sess = datetime.strptime(m.group(1), "%b %d, %Y").strftime("%Y-%m-%d")
                    for r in (blk.get("table") or {}).get("rows") or []:
                        pct = num(r.get("change"))
                        if pct is None:
                            continue
                        got.append(row(r["symbol"], r.get("name", ""), direction, pct, num(r.get("lastSalePrice")),
                                       num(r.get("lastSaleChange")), None, None, None, "nasdaq"))
            except Exception as e:
                lines.append(["nasdaq", ex, code, 0, "parse error " + str(e)[:60]])
                continue
        out += got
        lines.append(["nasdaq", ex, code, len(got), "" if got else "no rows"])
        time.sleep(random.uniform(2, 4))
    return out, sess, lines


def src_stockanalysis(now):
    out, lines = [], []
    for kind, direction in (("gainers", "Gainer"), ("losers", "Loser")):
        code, h = http_get(f"https://stockanalysis.com/markets/{kind}/")
        got = []
        if code == 200:
            for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", h, re.S):
                c = [text(x) for x in re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)]
                if len(c) < 7 or not re.fullmatch(r"[A-Z0-9.\-]{1,10}", c[1]):
                    continue
                pct = num(c[3])
                if pct is None:
                    continue
                got.append(row(c[1], c[2], direction, pct, num(c[4]), None, num(c[5]), None, num(c[6]), "stockanalysis"))
        out += got
        lines.append(["stockanalysis", kind, code, len(got), "" if got else "no rows parsed"])
        time.sleep(random.uniform(2, 4))
    return out, None, lines


FETCH = {"yahoo": src_yahoo, "nasdaq": src_nasdaq, "stockanalysis": src_stockanalysis}


# ----------------------------------------------------------------------------------------------- sectors
def load_sector_map():
    m = {}
    if not os.path.exists(SECTOR_MAP):
        print(f"WARNING: {os.path.basename(SECTOR_MAP)} not found - layer5_sector will be blank")
        return m
    for r in csv.DictReader(open(SECTOR_MAP, encoding="utf-8-sig")):
        m[norm(r["ticker"])] = r["sectors"]
    return m


class YahooProfiles:
    """Sector / industry from Yahoo's quoteSummary assetProfile (needs a cookie + crumb); falls back to the
    company-profile page if the crumb cannot be obtained."""

    def __init__(self):
        self.crumb = None
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
        self.opener.addheaders = [("User-Agent", UA), ("Accept-Language", "en-US,en;q=0.9")]
        try:
            try:
                self.opener.open("https://fc.yahoo.com", timeout=20)       # sets the cookie (answers 404, that is fine)
            except urllib.error.HTTPError:
                pass
            self.crumb = self.opener.open("https://query1.finance.yahoo.com/v1/test/getcrumb", timeout=20).read().decode()
            if not self.crumb or len(self.crumb) > 40 or "<" in self.crumb:
                self.crumb = None
        except Exception:
            self.crumb = None

    def lookup(self, sym):
        """Returns (http_code, sector, industry)."""
        if self.crumb:
            try:
                b = self.opener.open("https://query1.finance.yahoo.com/v10/finance/quoteSummary/"
                                     f"{urllib.parse.quote(sym)}?modules=assetProfile&crumb={urllib.parse.quote(self.crumb)}",
                                     timeout=20).read().decode("utf-8", "ignore")
                p = json.loads(b)["quoteSummary"]["result"][0].get("assetProfile") or {}
                return 200, p.get("sector") or "", p.get("industry") or ""
            except urllib.error.HTTPError as e:
                if e.code in (429, 503):
                    return e.code, "", ""
                if e.code == 404:                       # no asset profile (ETFs, some funds)
                    return 200, "", ""
            except Exception:
                pass
        code, h = http_get(f"https://finance.yahoo.com/quote/{sym}/profile/")
        if code != 200:
            return code, "", ""
        m = re.search(r"Sector:\s*(.+?)\s+Industry:\s*(.+?)\s+(?:Full Time|Employees|Description|Website)", text(h))
        return 200, (m.group(1) if m else ""), (m.group(2) if m else "")


# ----------------------------------------------------------------------------------------------- catalyst hints
def short_name(name):
    n = re.sub(r"\s*\(.*?\)", "", name or "")
    prev = None
    while prev != n:                                     # strip stacked suffixes ("Holdings, Inc.")
        prev, n = n, SUFFIX.sub("", n).strip(" ,.")
    return n or (name or "").strip()


def news_hints(sym, name, now, want=3):
    """Newest Google News headlines of the last 2 days that name the company or ticker. Returns (code, [dict])."""
    nm = short_name(name)
    terms = []
    if len(nm) >= 3:
        terms.append(f'"{nm}"')
    if len(sym) >= 3:
        terms.append(f'"{sym}"')
    if not terms:
        return 0, []
    q = "(" + " OR ".join(terms) + ") when:2d"
    code, xml = http_get(NEWS_URL.format(q=urllib.parse.quote(q, safe=":()")), timeout=30)
    if code != 200:
        return code, []
    cutoff = now - timedelta(hours=48)
    sym_re = re.compile(r"(?<![A-Z0-9])" + re.escape(sym) + r"(?![A-Z0-9])") if len(sym) >= 3 else None
    out, seen = [], set()
    for it in re.findall(r"<item>.*?</item>", xml, re.S):
        m = re.search(r"<title>(.*?)</title>", it, re.S)
        pd = re.search(r"<pubDate>(.*?)</pubDate>", it)
        if not (m and pd):
            continue
        title = html.unescape(re.sub(r"<!\[CDATA\[|\]\]>", "", m.group(1))).strip()
        try:
            dt = email.utils.parsedate_to_datetime(pd.group(1)).astimezone(ET)
        except Exception:
            continue
        s = re.search(r"<source[^>]*>(.*?)</source>", it, re.S)
        outlet = html.unescape(s.group(1)).strip() if s else ""
        head = re.sub(r"\s+-\s+" + re.escape(outlet) + r"$", "", title) if outlet else title
        if dt < cutoff or JUNK.search(head):
            continue
        if not ((len(nm) >= 3 and nm.lower() in head.lower()) or (sym_re and sym_re.search(head))):
            continue                                      # headline names neither the company nor the ticker
        key = re.sub(r"[^a-z0-9 ]", "", head.lower())[:80]
        if key in seen:
            continue
        seen.add(key)
        link = (re.search(r"<link>(.*?)</link>", it, re.S) or [None, ""])[1].strip()
        out.append({"dt": dt, "outlet": outlet, "headline": head, "link": link})
    out.sort(key=lambda x: x["dt"], reverse=True)
    return 200, out[:want]


def fill_hints(ranked, now):
    """Adds catalyst_hint_1..3 to each row. Stops early after two blocked requests in a row."""
    blocked = found = n_url = 0
    for i, r in enumerate(ranked, 1):
        if blocked >= 2:
            print("   hints stopped: two blocked requests in a row")
            break
        code, hints = news_hints(r["symbol"], r["name"], now)
        if code in (429, 503):
            blocked += 1
        else:
            blocked = 0
        for k, h in enumerate(hints, 1):
            url, st = gnews_resolve.resolve(h["link"])
            n_url += bool(url)
            r[f"catalyst_hint_{k}"] = " | ".join([h["dt"].strftime("%Y-%m-%d %H:%M"), h["outlet"], h["headline"],
                                                   url or h["link"]])
        found += bool(hints)
        if i % 20 == 0:
            print(f"   hints: {i}/{len(ranked)} movers, {found} with at least one headline")
        time.sleep(random.uniform(2, 4))
    return found, n_url


# ----------------------------------------------------------------------------------------------- workbook
def write_sheet(rows, session):
    wb = openpyxl.load_workbook(XLSX) if os.path.exists(XLSX) else openpyxl.Workbook()
    if wb.sheetnames == ["Sheet"]:
        del wb["Sheet"]
    if session in wb.sheetnames:
        del wb[session]
    ws = wb.create_sheet(session, 0)                       # newest sheet first
    ws.append(HEADER)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([r.get(h) for h in HEADER])
    for i, w in enumerate([6, 9, 38, 10, 11, 10, 9, 14, 14, 15, 44, 20, 28, 14, 14, 12, 18, 120, 120, 120], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    for rr in ws.iter_rows(min_row=2):
        rr[4].number_format = '0.00"%"'
        for k in (7, 8, 9):
            rr[k].number_format = "#,##0"
    path = XLSX
    try:
        wb.save(path)
    except PermissionError:
        path = os.path.join(HERE, f"market-movers.PENDING-{session}.xlsx")
        wb.save(path)
        print(f"  market-movers.xlsx is open in another program - saved to {os.path.basename(path)} instead")
    return path


# ----------------------------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--top", type=int, default=100)
    ap.add_argument("--source", choices=SOURCE_ORDER, help="first source to try (default yahoo)")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-profile", action="store_true")
    ap.add_argument("--no-hints", action="store_true", help="skip the catalyst-hint (news headline) lookup")
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    now = datetime.now(ET)
    order = SOURCE_ORDER[SOURCE_ORDER.index(a.source):] if a.source else SOURCE_ORDER
    pool, session, logs, used = {}, None, [], []
    for name in order:
        print(f"source: {name} ...")
        rows, sess, lines = FETCH[name](now)
        for ln in lines:
            logs.append([now.strftime("%Y-%m-%d %H:%M:%S")] + [sess or ""] + ln)
            print(f"   {ln[1]:<9} http {ln[2]}  rows {ln[3]}  {ln[4]}")
        added = 0
        for r in rows:
            if r["symbol"] not in pool:
                pool[r["symbol"]] = r
                added += 1
        if added:
            used.append(name)
        session = session or sess
        have = len(pool)
        print(f"   -> {added} new symbols, {have} total")
        if have >= a.top:
            break
        if name != order[-1]:
            print(f"   only {have} of {a.top} so far - falling back to the next source")

    if not pool:
        print("ERROR: no source returned any movers")
        log_rows(logs + [[now.strftime("%Y-%m-%d %H:%M:%S"), "", "none", "", "", 0, "ALL SOURCES FAILED"]])
        return 1

    if not session:                                        # source gave no date: last completed/ongoing session
        d = now.date()
        if now.weekday() < 5 and now.hour * 60 + now.minute < 9 * 60 + 30:
            d -= timedelta(days=1)
        while d.weekday() >= 5:
            d -= timedelta(days=1)
        session = d.strftime("%Y-%m-%d")
    ranked = sorted(pool.values(), key=lambda r: -abs(r["pct_change"]))[:a.top]
    smap = load_sector_map()
    stamp = now.strftime("%Y-%m-%d %H:%M")
    unmatched = []
    for i, r in enumerate(ranked, 1):
        r["rank"], r["session_date"], r["collected_et"] = i, session, stamp
        sec = smap.get(r["symbol"], "")
        r["layer5_sector"] = sec
        r["sector_source"] = "Layer 5 list" if sec else ""
        if not sec:
            unmatched.append(r)
    if not a.no_profile and unmatched and not a.dry_run:
        print(f"looking up Yahoo sector for {len(unmatched)} names outside Layer 5 ...")
        yp = YahooProfiles()
        print("   crumb lookup: " + ("ok" if yp.crumb else "unavailable, using profile pages"))
        for r in unmatched:
            code, s, ind = yp.lookup(r["symbol"])
            if code != 200:
                print(f"   profile lookup stopped at {r['symbol']} (http {code})")
                break
            r["yahoo_sector"], r["yahoo_industry"] = s, ind
            if s:
                r["sector_source"] = "Yahoo profile"
            time.sleep(random.uniform(0.5, 1.2))
    for r in ranked:
        if not r["sector_source"]:
            r["sector_source"] = "unclassified"

    hint_note = ""
    if not a.no_hints and not a.dry_run:
        print(f"looking up catalyst hints for {len(ranked)} movers ...")
        found, n_url = fill_hints(ranked, now)
        hint_note = f"; catalyst hints for {found} of {len(ranked)} movers ({n_url} article URLs resolved)"
        print(f"   -> {found} movers have at least one headline")

    n_l5 = sum(1 for r in ranked if r["layer5_sector"])
    n_y = sum(1 for r in ranked if r["sector_source"] == "Yahoo profile")
    print(f"\nsession {session}: {len(ranked)} movers from {', '.join(used)} | "
          f"{n_l5} in Layer 5 sectors, {n_y} sector from Yahoo profile, {len(ranked) - n_l5 - n_y} unclassified")
    for r in ranked[:10]:
        print(f"  {r['rank']:>3} {r['symbol']:<7} {r['pct_change']:>8.2f}%  {r['name'][:30]:<30} "
              f"{r['layer5_sector'] or r.get('yahoo_sector', '')}")
    if a.dry_run:
        print("(dry run - nothing written)")
        return 0
    path = write_sheet(ranked, session)
    note = f"{len(ranked)} rows written; sources used: {', '.join(used)}{hint_note}"
    log_rows(logs + [[now.strftime("%Y-%m-%d %H:%M:%S"), session, "+".join(used), "TOTAL", "", len(ranked), note]])
    print(f"wrote sheet {session} -> {os.path.basename(path)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
